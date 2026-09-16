"""
Computes each household income quintile's share of aggregate Social Security
benefits, using CBO's "The Distribution of Household Income, 2022" (Jan 2026
release, www.cbo.gov/publication/61911), Supplemental Data workbook.

Source file: data/cbo_61911_supplemental_data.xlsx
  - Table 1 ("1. Demographics"): number of households by income quintile, 2022
  - Table 5 ("5. Income Before Trans + Tax"): average Social Security benefit
    per household, by income quintile, 2022 (2022 dollars)

Methodology, per CBO's own notes (Contents and Notes sheet):
  Households are ranked into quintiles by "income before transfers and taxes"
  (market income plus Social Security, Medicare, unemployment insurance, and
  workers' compensation), adjusted for household size. Each quintile contains
  approximately equal numbers of PEOPLE, not households, so household counts
  differ slightly across quintiles -- which is why this script weights each
  quintile's average benefit by its own household count rather than assuming
  equal-sized quintiles.

  Note this ranking measure already includes Social Security benefits
  themselves (via "social insurance benefits"), not pure market income alone.
  CBO does not publish a standard quintile breakdown ranked on market income
  excluding social insurance benefits.

Aggregate share for quintile i = (avg SS benefit_i * households_i) / sum_j(avg SS benefit_j * households_j)
"""
import openpyxl
from pathlib import Path

SRC = Path(__file__).resolve().parent.parent / "data" / "cbo_61911_supplemental_data.xlsx"
YEAR = 2022
QUINTILES = ["Lowest quintile", "Second quintile", "Middle quintile", "Fourth quintile", "Highest quintile"]

wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)

# Table 1: households (millions) by quintile
demo_rows = list(wb["1. Demographics"].iter_rows(min_row=9, values_only=True))
households = {r[0]: r[2] for r in demo_rows[1:] if r[1] == YEAR}

# Table 5: avg Social Security benefit per household ($, 2022 dollars) by quintile
inc_rows = list(wb["5. Income Before Trans + Tax"].iter_rows(min_row=9, values_only=True))
avg_ss = {r[0]: r[19] for r in inc_rows[1:] if r[1] == YEAR}  # column 19 = "Social Security benefits"

aggregate = {q: avg_ss[q] * households[q] for q in QUINTILES}  # $ millions
total = sum(aggregate.values())

print(f"{'Quintile':<18}{'Avg SS/HH ($)':>15}{'Households (M)':>16}{'Aggregate ($M)':>17}{'Share':>9}")
for q in QUINTILES:
    print(f"{q:<18}{avg_ss[q]:>15,.0f}{households[q]:>16.1f}{aggregate[q]:>17,.0f}{aggregate[q]/total*100:>8.1f}%")
print(f"{'Total':<18}{'':>15}{'':>16}{total:>17,.0f}{100.0:>8.1f}%")
print(f"\nImplied total Social Security benefits received by households, {YEAR}: ${total/1e6:,.2f} trillion")
